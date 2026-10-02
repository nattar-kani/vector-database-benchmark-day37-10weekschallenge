from pathlib import Path
from docx import Document as DocxDocument
from bs4 import BeautifulSoup
import pandas as pd
import fitz
import pdfplumber

from .models import Document

def load_txt(
        file_path: str | Path
) -> Document:
    path = Path(file_path)
    content = path.read_text(encoding="utf-8")

    return Document(
        content=content,
        source=str(path),
        file_type="txt",
        metadata={
            "file_name":path.name,
            "file_size": path.stat().st_size
        }
    )

def load_docx(
        file_path: str | Path
) -> Document:
    path = Path(file_path)
    docx = DocxDocument(path)

    parts = []

    for paragraphs in docx.paragraphs:
        text = paragraphs.text.strip()

        if text:
            parts.append(text)

    for table in docx.tables:
        for row in table.rows:
            row_data = [cell.text.strip() for cell in row.cells]
            parts.append("|".join(row_data))

    content = "\n".join(parts)

    return Document(
        content=content,
        source=str(path),
        file_type="docx",
        metadata={
            "file_name": path.name,
            "paragraph_count": len(docx.paragraphs),
            "table_count": len(docx.tables),
            "file_size": path.stat().st_size
        }
    )

def load_html(
        file_path: str | Path
) -> Document:
    path = Path(file_path)
    html = path.read_text(encoding="utf-8")

    soup = BeautifulSoup(html, "html.parser")

    for  element in soup(["script","style"]):
        element.decompose()

    content = soup.get_text(separator="\n",strip=True)

    return Document(
        content=content,
        source=str(path),
        file_type="html",
        metadata={
            "file_name": path.name,
            "title": soup.title.string.strip() if soup.title else None,
            "file_size": path.stat().st_size
        }
    )

def load_csv(
        file_path: str | Path
) -> Document:
    path = Path(file_path)
    df = pd.read_csv(path)

    rows = []

    rows.append("|".join(df.columns.astype(str)))

    for _,row in df.iterrows():
        rows.append("|".join(row.astype(str)))

    content = "\n".join(rows)

    return Document(
        content=content,
        source=str(path),
        file_type="csv",
        metadata={
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "row_count": len(df),
            "column_count": len(df.columns),
            "file_size": path.stat().st_size
        }
    )

def load_pdf(
    file_path: str | Path
)-> Document:
    path = Path(file_path)

    pdf = fitz.open(path)

    pages =[]

    image_count = 0

    for page in pdf:
        text = page.get_text().strip()

        if text:
            pages.append(text)

        images = page.get_images(full=True)
        image_count += len(images)

    tables = []

    with pdfplumber.open(path) as pdf_plumber:
        for page in pdf_plumber.pages:
            page_tables = page.extract_tables()

            for table in page_tables:
                for row in table:
                    row_data = [
                        str(cell).strip() if cell is not None else ""
                        for cell in row]

                    tables.append(" | ".join(row_data))


    content_parts = pages + tables

    content = "\n\n".join(content_parts)

    return Document(
        content=content,
        source=str(path),
        file_type="pdf",
        metadata={
            "file_name": path.name,
            "file_size": path.stat().st_size,
            "page_count": len(pdf),
            "table_count": len(tables),
            "image_count": image_count,
            "file_size": path.stat().st_size
        },
    )

def load_document(
        file_path:str | Path
) -> Document:
    
    path = Path(file_path)
    
    loaders = {
        ".txt": load_txt,
        ".docx": load_docx,
        ".html": load_html,
        ".csv": load_csv,
        ".pdf": load_pdf,
    }

    loader = loaders.get(path.suffix.lower())

    if loader is None:
        raise ValueError(
            f"Unsupported file type: {path.suffix}"
        )

    return loader(path)

                    


















